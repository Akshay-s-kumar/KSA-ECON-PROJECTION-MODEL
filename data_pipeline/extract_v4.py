"""Version 4 extraction: Excel -> v4_reference.duckdb (run once per workbook change).

Reads the raw comparator regions and the Excel checkpoint values. No scenario is precomputed.

    python extract_v4.py
    python extract_v4.py --workbook "path/to/workbook.xlsx" --output v4_reference.duckdb
"""

from __future__ import annotations

import argparse
import datetime as dt
import time
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string as col_idx

import v4_config as cfg


def _num(value) -> float:
    """Numbers stay numbers; text, blanks and errors become NaN (excluded from every filter)."""
    if isinstance(value, bool):
        return np.nan
    return float(value) if isinstance(value, (int, float)) else np.nan


def _block(ws, first_row: int, last_row: int, first_col: str, last_col: str) -> list[tuple]:
    return list(ws.iter_rows(min_row=first_row, max_row=last_row,
                             min_col=col_idx(first_col), max_col=col_idx(last_col),
                             values_only=True))


def _cell(ws, address: str):
    return _block(ws, int("".join(filter(str.isdigit, address))),
                  int("".join(filter(str.isdigit, address))),
                  "".join(filter(str.isalpha, address)),
                  "".join(filter(str.isalpha, address)))[0][0]


def _band(ws, cells: dict, lookup: str, sector_code: str | None) -> dict:
    band = {key: _num(_cell(ws, address)) for key, address in cells.items()}
    missing = [key for key, value in band.items() if np.isnan(value)]
    if missing:
        raise ValueError(f"{ws.title}: band cells {missing} are not numeric")
    return {"lookup": lookup, "sector_code": sector_code, **band}


def _column(rows: list[tuple], offset: int) -> list:
    return [row[offset] for row in rows]


def read_workbook(path: Path) -> dict[str, pd.DataFrame]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        gva_frames, gva_emp_frames, bands, checks = [], [], [], []

        # ---- GVA share sheets ------------------------------------------------
        g = cfg.GVA
        share_cols = range(col_idx(g["share_first_col"]), col_idx(g["share_last_col"]) + 1)
        for code, (gva_sheet, gva_emp_sheet) in cfg.SECTORS.items():
            ws = wb[gva_sheet]
            years = [int(y) for y in _block(ws, 1, 1, g["share_first_col"], g["share_last_col"])[0]]
            rows = _block(ws, g["first_row"], g["last_row"], "A", g["nuts2_pop_col"])
            base = col_idx("A")
            frame = pd.DataFrame({
                "sector_code": code,
                "excel_row": range(g["first_row"], g["last_row"] + 1),
                "nuts3_id": _column(rows, col_idx(g["id_col"]) - base),
                "nuts3_pop": [_num(v) for v in _column(rows, col_idx(g["nuts3_pop_col"]) - base)],
                "nuts2_pop": [_num(v) for v in _column(rows, col_idx(g["nuts2_pop_col"]) - base)],
            })
            for year, column in zip(years, share_cols):
                frame[f"share_{year}"] = [_num(v) for v in _column(rows, column - base)]
            gva_frames.append(frame)
            bands.append(_band(ws, g["band_cells"], "gva", code))

            check_rows = _block(ws, g["check_first_row"], g["check_last_row"],
                                g["check_percentile_col"], g["check_value_col"])
            p_off = 0
            v_off = col_idx(g["check_value_col"]) - col_idx(g["check_percentile_col"])
            for row in check_rows:
                checks.append({"lookup": "gva", "sector_code": code,
                               "percentile": round(_num(row[p_off]), 4), "excel_value": _num(row[v_off])})

            # ---- GVA-employment delta sheet ---------------------------------
            e = cfg.GVA_EMP
            ws = wb[gva_emp_sheet]
            rows = _block(ws, e["first_row"], e["last_row"], "A", e["nuts2_pop_col"])
            gva_emp_frames.append(pd.DataFrame({
                "sector_code": code,
                "excel_row": range(e["first_row"], e["last_row"] + 1),
                "nuts3_id": _column(rows, col_idx(e["id_col"]) - base),
                "nuts3_pop": [_num(v) for v in _column(rows, col_idx(e["nuts3_pop_col"]) - base)],
                "nuts2_pop": [_num(v) for v in _column(rows, col_idx(e["nuts2_pop_col"]) - base)],
                "value": [_num(v) for v in _column(rows, col_idx(e["value_col"]) - base)],
            }))
            bands.append(_band(ws, e["band_cells"], "gva_emp", code))

        # ---- Employment-population delta (one sheet) ---------------------------
        p = cfg.EMP_POP
        ws = wb[p["sheet"]]
        rows = _block(ws, p["first_row"], p["last_row"], "A", p["nuts2_pop_col"])
        base = col_idx("A")
        emp_pop = pd.DataFrame({
            "excel_row": range(p["first_row"], p["last_row"] + 1),
            "nuts3_id": _column(rows, col_idx(p["id_col"]) - base),
            "nuts3_pop": [_num(v) for v in _column(rows, col_idx(p["nuts3_pop_col"]) - base)],
            "nuts2_pop": [_num(v) for v in _column(rows, col_idx(p["nuts2_pop_col"]) - base)],
            "value": [_num(v) for v in _column(rows, col_idx(p["value_col"]) - base)],
        })
        bands.append(_band(ws, p["band_cells"], "emp_pop", None))
        v_off = col_idx(p["check_value_col"]) - col_idx(p["check_percentile_col"])
        for row in _block(ws, p["check_first_row"], p["check_last_row"],
                          p["check_percentile_col"], p["check_value_col"]):
            pct, value = _num(row[0]), _num(row[v_off])
            if not np.isnan(pct) and not np.isnan(value):
                checks.append({"lookup": "emp_pop", "sector_code": None,
                               "percentile": round(pct, 4), "excel_value": value})
    finally:
        wb.close()

    return {
        "gva_comparators": pd.concat(gva_frames, ignore_index=True),
        "gva_emp_comparators": pd.concat(gva_emp_frames, ignore_index=True),
        "emp_pop_comparators": emp_pop,
        "default_bands": pd.DataFrame(bands),
        "excel_checkpoints": pd.DataFrame(checks),
        "metadata": pd.DataFrame([{
            "workbook": path.name,
            "workbook_modified": dt.datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds"),
            "extracted_at": dt.datetime.now().isoformat(timespec="seconds"),
        }]),
    }


def write_duckdb(tables: dict[str, pd.DataFrame], output: Path) -> None:
    import duckdb

    temp = output.with_suffix(".tmp.duckdb")
    temp.unlink(missing_ok=True)
    con = duckdb.connect(str(temp))
    try:
        for name, frame in tables.items():
            con.register("frame", frame)
            con.execute(f"CREATE TABLE {name} AS SELECT * FROM frame")
            con.unregister("frame")
    finally:
        con.close()
    temp.replace(output)          # fails cleanly if the API still has the old file open


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--workbook", type=Path, default=cfg.WORKBOOK_PATH)
    parser.add_argument("--output", type=Path, default=cfg.REFERENCE_DB_PATH)
    args = parser.parse_args()

    start = time.perf_counter()
    print(f"Reading {args.workbook.name} (large workbook, this can take a minute)...")
    tables = read_workbook(args.workbook)
    write_duckdb(tables, args.output)
    for name, frame in tables.items():
        print(f"  {name:<22} {len(frame):>7,} rows")
    print(f"Saved {args.output} in {time.perf_counter() - start:.1f}s")


if __name__ == "__main__":
    main()
