"""Excel-free Version 3 scenario calculations.

The workbook is used once as a reference dataset. Runtime scenario generation
uses the extracted lookup tables and the formulas confirmed in the workbook.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import UTC, datetime
from itertools import product
from math import exp, log
from pathlib import Path
from typing import Iterable

import duckdb
import numpy as np
from openpyxl import load_workbook


WORKSHEET = "GVA-EmploymentShares&Projec_HYB"
WORKBOOK_VERSION = "M14"
REGION = "Riyadh"
START_YEAR = 2023
TARGET_YEAR = 2050
DEFAULT_STEEPNESS = 0.25
BASE_VALUE = 0.0001
GVA_PERCENTILES = tuple(round(index / 20, 2) for index in range(1, 21))
DELTA_PERCENTILES = tuple(round(index / 20, 2) for index in range(21))

SECTOR_COLUMN_MAP = {
    "A": "C",
    "B-E": "D",
    "F": "E",
    "G-I": "F",
    "J": "G",
    "K": "H",
    "L": "I",
    "M_N": "J",
    "O-Q": "K",
    "R-U": "L",
}

SECTOR_NAMES = {
    "A": "Agriculture, forestry and fishing",
    "B-E": "Manufacturing, mining, quarrying and other industry",
    "F": "Construction",
    "G-I": "Wholesale and retail trade, transportation and storage, accommodation and food service activities",
    "J": "Information and communication",
    "K": "Financial and insurance activities",
    "L": "Real estate activities",
    "M_N": "Professional, scientific, technical, administrative and support service activities",
    "O-Q": "Public administration, defence, education, human health and social work activities",
    "R-U": "Other services",
}

SECTOR_ORDER = tuple(SECTOR_COLUMN_MAP)
SECTOR_INDEX = {sector: index for index, sector in enumerate(SECTOR_ORDER)}

GVA_SHEETS = {
    "A": "NUTS3-A",
    "B-E": "NUTS3-B-E",
    "F": "NUTS3-F",
    "G-I": "NUTS3-G-I",
    "J": "NUTS3-J",
    "K": "NUTS3-K",
    "L": "NUTS3-L",
    "M_N": "NUTS3-M_N",
    "O-Q": "NUTS3-O-Q",
    "R-U": "NUTS3-R-U",
}

GVA_EMP_SHEETS = {
    "A": "GVA_EMP_DELTA-3-A",
    "B-E": "GVA_EMP_DELTA-3-B-E",
    "F": "GVA_EMP_DELTA-3-F",
    "G-I": "GVA_EMP_DELTA-3-G-I",
    "J": "GVA_EMP_DELTA-3-J",
    "K": "GVA_EMP_DELTA-3-K",
    "L": "GVA_EMP_DELTA-3-L",
    "M_N": "GVA_EMP_DELTA-3-M-N",
    "O-Q": "GVA_EMP_DELTA-3-O-Q",
    "R-U": "GVA_EMP_DELTA-3-R-U",
}


def excel_percentile(values: Iterable[float], percentile: float) -> float:
    """Match Excel's inclusive PERCENTILE interpolation for numeric values."""

    array = np.asarray(tuple(values), dtype=float)
    if array.size == 0:
        raise ValueError("Cannot calculate a percentile from an empty series.")
    if not 0 <= percentile <= 1:
        raise ValueError(f"Percentile must be between 0 and 1: {percentile}")
    return float(np.percentile(array, percentile * 100, method="linear"))


def calculate_inflection_point(
    target_value: float,
    base_value: float = BASE_VALUE,
    start_year: int = START_YEAR,
    steepness: float = DEFAULT_STEEPNESS,
) -> float:
    """Implement the workbook's row-11 inflection-point formula."""

    if target_value <= 0 or base_value <= 0:
        raise ValueError("Target and base values must be positive.")
    if not 0 < base_value < target_value:
        raise ValueError(
            "The workbook logistic curve requires base_value < target_value."
        )
    if steepness <= 0:
        raise ValueError("Steepness must be positive.")
    ratio = base_value / target_value
    return start_year + log((1 - ratio) / ratio) / steepness


def calculate_gva_share(
    year: int,
    target_value: float,
    base_value: float = BASE_VALUE,
    start_year: int = START_YEAR,
    steepness: float = DEFAULT_STEEPNESS,
) -> float:
    """Implement the workbook's annual GVA Share S-curve formula."""

    inflection = calculate_inflection_point(
        target_value=target_value,
        base_value=base_value,
        start_year=start_year,
        steepness=steepness,
    )
    return target_value / (1 + exp(-steepness * (year - inflection)))


@dataclass(frozen=True)
class Scenario:
    scenario_id: str
    region: str
    sector_code: str
    sector_name: str
    workbook_column: str
    gva_percentile: float
    gva_share_result: float
    gva_emp_delta_percentile: float
    gva_emp_delta_result: float
    emp_pop_delta_percentile: float
    emp_pop_delta_result: float
    workbook_version: str = WORKBOOK_VERSION


@dataclass(frozen=True)
class AnnualResult:
    scenario_id: str
    year: int
    gva_share_s_curve: float
    gva_sar_million: float
    employment_saudi: float
    employment_non_saudi: float
    employment_total: float
    population: float
    income: float
    consumption: float
    sector_gva_sar_million: float
    sector_employment_saudi: float
    sector_employment_non_saudi: float
    sector_employment_total: float
    sector_population: float
    sector_income: float
    sector_consumption: float


class WorkbookReference:
    """Extract the confirmed lookup tables from the workbook once."""

    def __init__(self, workbook_path: Path) -> None:
        self.workbook_path = workbook_path
        workbook = load_workbook(workbook_path, data_only=True, read_only=True)
        try:
            self.gva_lookup = self._read_gva_lookup(workbook)
            self.gva_emp_lookup = self._read_gva_emp_lookup(workbook)
            emp_pop_sheet = workbook["Emp-Pop-Delta-3"]
            self.emp_pop_values = tuple(
                float(value)
                for (value,) in emp_pop_sheet.iter_rows(
                    min_row=2, max_row=1351, min_col=38, max_col=38,
                    values_only=True,
                )
                if isinstance(value, (int, float))
            )
            self.emp_pop_lookup = {
                percentile: excel_percentile(self.emp_pop_values, percentile)
                for percentile in DELTA_PERCENTILES
            }
            main_sheet = workbook[WORKSHEET]
            self.baseline_gva_percentiles = tuple(
                round(float(main_sheet.cell(row=5, column=column).value), 2)
                for column in range(3, 13)
            )
            self.baseline_gva_emp_percentiles = tuple(
                round(float(main_sheet.cell(row=13, column=column).value), 2)
                for column in range(3, 13)
            )
            self.baseline_emp_pop_percentile = float(
                main_sheet["C16"].value
            )
            source = workbook["Riyadh_GVA_Emp_NACE Sector"]
            self.base_gva = {
                int(source.cell(row=row, column=1).value): tuple(
                    float(source.cell(row=row, column=column).value)
                    for column in range(2, 12)
                )
                for row in range(9, 57)
            }
            self.productivity = {
                int(source.cell(row=row, column=1).value): tuple(
                    float(source.cell(row=row, column=column).value)
                    for column in range(2, 12)
                )
                for row in range(122, 170)
            }
            riyadh_oe = workbook["Riyadh_OE"]
            self.population_ratio = {
                int(riyadh_oe.cell(row=row, column=1).value): float(
                    riyadh_oe.cell(row=row, column=5).value
                ) / float(riyadh_oe.cell(row=row, column=2).value)
                for row in range(7, 55)
            }
            workforce = workbook["Saudi-Non Saudi Assumption"]
            self.workforce_share = {
                "Skilled Sectors": (
                    float(workforce["C112"].value),
                    float(workforce["F112"].value),
                ),
                "Mixed Sectors": (
                    float(workforce["D112"].value),
                    float(workforce["G112"].value),
                ),
                "Labour-Intensive sectors": (
                    float(workforce["E112"].value),
                    float(workforce["H112"].value),
                ),
            }
            income = workbook["Income Distribution Assumption"]
            self.income_rate = {
                "Skilled Sectors": (
                    float(income["B52"].value),
                    float(income["E52"].value),
                ),
                "Mixed Sectors": (
                    float(income["C52"].value),
                    float(income["F52"].value),
                ),
                "Labour-Intensive sectors": (
                    float(income["D52"].value),
                    float(income["G52"].value),
                ),
            }
            self.mpc = tuple(float(income.cell(row=76, column=column).value)
                            for column in range(2, 5))
            self.employment_groups = (
                "Labour-Intensive sectors",
                "Mixed Sectors",
                "Labour-Intensive sectors",
                "Mixed Sectors",
                "Skilled Sectors",
                "Skilled Sectors",
                "Mixed Sectors",
                "Skilled Sectors",
                "Mixed Sectors",
                "Mixed Sectors",
            )
        finally:
            workbook.close()

    @staticmethod
    def _read_gva_lookup(workbook) -> dict[str, dict[float, float]]:
        lookups: dict[str, dict[float, float]] = {}
        for sector_code, sheet_name in GVA_SHEETS.items():
            sheet = workbook[sheet_name]
            lookup: dict[float, float] = {}
            for percentile, _, value in sheet.iter_rows(
                min_row=1353, max_row=1372, min_col=58, max_col=60,
                values_only=True,
            ):
                if isinstance(percentile, (int, float)) and isinstance(
                    value, (int, float)
                ):
                    lookup[round(float(percentile), 2)] = float(value)
            if not lookup:
                raise ValueError(f"No GVA lookup values found in {sheet_name}.")
            lookups[sector_code] = lookup
        return lookups

    @staticmethod
    def _read_gva_emp_lookup(workbook) -> dict[str, dict[float, float]]:
        lookups: dict[str, dict[float, float]] = {}
        for sector_code, sheet_name in GVA_EMP_SHEETS.items():
            sheet = workbook[sheet_name]
            lookup: dict[float, float] = {}
            for column in range(40, sheet.max_column + 1):
                percentile = sheet.cell(row=2, column=column).value
                value = sheet.cell(row=3, column=column).value
                if isinstance(percentile, (int, float)) and isinstance(
                    value, (int, float)
                ) and 0 <= percentile <= 1:
                    lookup[round(float(percentile), 2)] = float(value)
            if not lookup:
                raise ValueError(
                    f"No GVA-employment lookup values found in {sheet_name}."
                )
            lookups[sector_code] = lookup
        return lookups

    def scenario(
        self,
        sector_code: str,
        gva_percentile: float,
        gva_emp_delta_percentile: float,
        emp_pop_delta_percentile: float,
    ) -> Scenario:
        if sector_code not in SECTOR_COLUMN_MAP:
            raise ValueError(f"Unsupported sector code: {sector_code}")
        values = (
            gva_percentile,
            gva_emp_delta_percentile,
            emp_pop_delta_percentile,
        )
        if gva_percentile not in GVA_PERCENTILES:
            raise ValueError(
                f"GVA percentile must be one of {GVA_PERCENTILES}."
            )
        if (
            gva_emp_delta_percentile not in DELTA_PERCENTILES
            or emp_pop_delta_percentile not in DELTA_PERCENTILES
        ):
            raise ValueError(
                f"Delta percentiles must be one of {DELTA_PERCENTILES}."
            )

        gva_result = self.gva_lookup[sector_code][gva_percentile]
        gva_emp_result = self.gva_emp_lookup[sector_code][
            gva_emp_delta_percentile
        ]
        emp_pop_result = self.emp_pop_lookup[emp_pop_delta_percentile]
        scenario_id = (
            f"{WORKBOOK_VERSION}_RIY_{sector_code.replace('-', '')}"
            f"_GVA{round(gva_percentile * 100):03d}"
            f"_GE{round(gva_emp_delta_percentile * 100):03d}"
            f"_EP{round(emp_pop_delta_percentile * 100):03d}"
        )
        return Scenario(
            scenario_id=scenario_id,
            region=REGION,
            sector_code=sector_code,
            sector_name=SECTOR_NAMES[sector_code],
            workbook_column=SECTOR_COLUMN_MAP[sector_code],
            gva_percentile=gva_percentile,
            gva_share_result=gva_result,
            gva_emp_delta_percentile=gva_emp_delta_percentile,
            gva_emp_delta_result=gva_emp_result,
            emp_pop_delta_percentile=emp_pop_delta_percentile,
            emp_pop_delta_result=emp_pop_result,
        )

    def annual_results(
        self,
        scenario: Scenario,
        years: Iterable[int],
    ) -> list[AnnualResult]:
        gva_percentiles = list(self.baseline_gva_percentiles)
        gva_emp_percentiles = list(self.baseline_gva_emp_percentiles)
        emp_pop_percentile = self.baseline_emp_pop_percentile
        selected_index = SECTOR_INDEX[scenario.sector_code]
        gva_percentiles[selected_index] = scenario.gva_percentile
        gva_emp_percentiles[selected_index] = scenario.gva_emp_delta_percentile
        emp_pop_percentile = scenario.emp_pop_delta_percentile

        targets = [
            self.gva_lookup[sector][percentile]
            for sector, percentile in zip(
                SECTOR_ORDER, gva_percentiles, strict=True
            )
        ]
        deltas = [
            self.gva_emp_lookup[sector][percentile]
            for sector, percentile in zip(
                SECTOR_ORDER, gva_emp_percentiles, strict=True
            )
        ]
        population_multiplier = self.emp_pop_lookup[emp_pop_percentile]
        results = []
        for year in years:
            if year not in self.base_gva or year not in self.productivity:
                raise ValueError(f"Unsupported projection year: {year}")
            shares = [
                calculate_gva_share(year, target)
                for target in targets
            ]
            gva_values = [
                share * base
                for share, base in zip(
                    shares, self.base_gva[year], strict=True
                )
            ]
            employment = [
                gva * 1_000_000 / productivity / (1 + delta)
                for gva, productivity, delta in zip(
                    gva_values, self.productivity[year], deltas, strict=True
                )
            ]
            employment_saudi = 0.0
            employment_non_saudi = 0.0
            income_total = 0.0
            sector_employment_saudi = 0.0
            sector_employment_non_saudi = 0.0
            sector_income = 0.0
            for index, (jobs, group) in enumerate(
                zip(employment, self.employment_groups, strict=True)
            ):
                saudi_share, non_saudi_share = self.workforce_share[group]
                employment_saudi += jobs * saudi_share
                employment_non_saudi += jobs * non_saudi_share
                saudi_rate, non_saudi_rate = self.income_rate[group]
                income = (
                    12 * jobs * saudi_share * saudi_rate / 1_000_000
                    + 12 * jobs * non_saudi_share * non_saudi_rate / 1_000_000
                )
                income_total += income
                if index == selected_index:
                    sector_employment_saudi = jobs * saudi_share
                    sector_employment_non_saudi = jobs * non_saudi_share
                    sector_income = income
            employment_total = sum(employment)
            population = (
                employment_total
                / self.population_ratio[year]
                / (1 + population_multiplier)
            )
            consumption = income_total * 1_000_000 * self.mpc[2]
            sector_employment_total = employment[selected_index]
            sector_population = (
                sector_employment_total
                / self.population_ratio[year]
                / (1 + population_multiplier)
            )
            sector_consumption = sector_income * 1_000_000 * self.mpc[2]
            results.append(
                AnnualResult(
                    scenario_id=scenario.scenario_id,
                    year=year,
                    gva_share_s_curve=shares[selected_index],
                    gva_sar_million=sum(gva_values),
                    employment_saudi=employment_saudi,
                    employment_non_saudi=employment_non_saudi,
                    employment_total=employment_total,
                    population=population,
                    income=income_total,
                    consumption=consumption,
                    sector_gva_sar_million=gva_values[selected_index],
                    sector_employment_saudi=sector_employment_saudi,
                    sector_employment_non_saudi=sector_employment_non_saudi,
                    sector_employment_total=sector_employment_total,
                    sector_population=sector_population,
                    sector_income=sector_income,
                    sector_consumption=sector_consumption,
                )
            )
        return results


def annual_results(
    scenario: Scenario,
    years: Iterable[int],
) -> list[AnnualResult]:
    raise RuntimeError(
        "Use WorkbookReference.annual_results so dependent workbook inputs "
        "are applied to every sector."
    )


def generate_duckdb(
    reference: WorkbookReference,
    database_path: Path,
    years: Iterable[int],
    sector_codes: Iterable[str] | None = None,
) -> int:
    """Generate all 88,200 confirmed-output scenarios into DuckDB."""

    year_values = tuple(years)
    if not year_values:
        raise ValueError("At least one projection year is required.")
    generated_at = datetime.now(UTC)
    selected_sectors = tuple(sector_codes or SECTOR_COLUMN_MAP)
    invalid_sectors = set(selected_sectors) - set(SECTOR_COLUMN_MAP)
    if invalid_sectors:
        raise ValueError(f"Unsupported sector codes: {sorted(invalid_sectors)}")

    connection = duckdb.connect(str(database_path))
    try:
        connection.execute("DROP TABLE IF EXISTS scenario_master")
        connection.execute("DROP TABLE IF EXISTS scenario_annual_results")
        connection.execute(
            """
            CREATE TABLE scenario_master (
                scenario_id VARCHAR PRIMARY KEY,
                workbook_version VARCHAR,
                region VARCHAR,
                sector_code VARCHAR,
                sector_name VARCHAR,
                workbook_column VARCHAR,
                gva_percentile DOUBLE,
                gva_share_result DOUBLE,
                gva_emp_delta_percentile DOUBLE,
                gva_emp_delta_result DOUBLE,
                emp_pop_delta_percentile DOUBLE,
                emp_pop_delta_result DOUBLE,
                generated_at TIMESTAMP
            )
            """
        )
        connection.execute(
            """
            CREATE TABLE scenario_annual_results (
                scenario_id VARCHAR,
                year INTEGER,
                gva_share_s_curve DOUBLE,
                gva_sar_million DOUBLE,
                employment_saudi DOUBLE,
                employment_non_saudi DOUBLE,
                employment_total DOUBLE,
                population DOUBLE,
                income DOUBLE,
                consumption DOUBLE,
                sector_gva_sar_million DOUBLE,
                sector_employment_saudi DOUBLE,
                sector_employment_non_saudi DOUBLE,
                sector_employment_total DOUBLE,
                sector_population DOUBLE,
                sector_income DOUBLE,
                sector_consumption DOUBLE,
                PRIMARY KEY (scenario_id, year)
            )
            """
        )
        master_rows = []
        annual_rows = []
        scenario_count = 0
        for sector_code in selected_sectors:
            for gva, gva_emp, emp_pop in product(
                GVA_PERCENTILES,
                DELTA_PERCENTILES,
                DELTA_PERCENTILES,
            ):
                scenario = reference.scenario(sector_code, gva, gva_emp, emp_pop)
                master_rows.append(
                    (
                        scenario.scenario_id,
                        scenario.workbook_version,
                        scenario.region,
                        scenario.sector_code,
                        scenario.sector_name,
                        scenario.workbook_column,
                        scenario.gva_percentile,
                        scenario.gva_share_result,
                        scenario.gva_emp_delta_percentile,
                        scenario.gva_emp_delta_result,
                        scenario.emp_pop_delta_percentile,
                        scenario.emp_pop_delta_result,
                        generated_at,
                    )
                )
                annual_rows.extend(
                    (
                        item.scenario_id,
                        item.year,
                        item.gva_share_s_curve,
                        item.gva_sar_million,
                        item.employment_saudi,
                        item.employment_non_saudi,
                        item.employment_total,
                        item.population,
                        item.income,
                        item.consumption,
                        item.sector_gva_sar_million,
                        item.sector_employment_saudi,
                        item.sector_employment_non_saudi,
                        item.sector_employment_total,
                        item.sector_population,
                        item.sector_income,
                        item.sector_consumption,
                    )
                    for item in reference.annual_results(scenario, year_values)
                )
                scenario_count += 1
                if len(master_rows) >= 10000:
                    connection.executemany(
                        "INSERT INTO scenario_master VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                        master_rows,
                    )
                    connection.executemany(
                        "INSERT INTO scenario_annual_results VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                        annual_rows,
                    )
                    master_rows.clear()
                    annual_rows.clear()
        if master_rows:
            connection.executemany(
                "INSERT INTO scenario_master VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                master_rows,
            )
            connection.executemany(
                "INSERT INTO scenario_annual_results VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                annual_rows,
            )
        connection.commit()
    finally:
        connection.close()
    return scenario_count
